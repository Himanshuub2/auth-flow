"""Combined events + documents: list, detail, revisions, snapshot, KPI, filter."""

import io
import logging
from datetime import date, datetime
from decimal import Decimal

from fastapi import APIRouter, Body, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from constants import DOCUMENT, EVENT
from database import get_db
from schemas.documents.combined import CombinedItemOut
from schemas.documents.items_filter import ItemsListBody
from schemas.events.comman import APIResponse, APIResponsePaginated
from services import items_service
from utils.dates import format_date_dmy_month_abbr
from utils.security import CurrentUser, get_current_user
from sqlalchemy.ext.asyncio import AsyncSession

router = APIRouter()
logger = logging.getLogger(__name__)

EXPORT_COLUMNS = (
    ("id", "ID"),
    ("name", "Name"),
    ("document_type", "Document Type"),
    ("status", "Status"),
    ("created_by_name", "Created By Name"),
    ("updated_at", "Updated At"),
    ("deactivated_by", "Deactivated By"),
    ("deactivated_by_name", "Deactivated By Name"),
    ("deactivated_at", "Deactivated At"),
    ("next_review_date", "Next Review Date"),
    ("revision", "Revision"),
    ("version", "Version"),
)
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _excel_cell_value(value):
    if isinstance(value, (datetime, date)):
        return format_date_dmy_month_abbr(value)
    if isinstance(value, Decimal):
        return float(value)
    return value


def _build_combined_export_workbook(items: list[CombinedItemOut]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Combined Items"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col_idx, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_idx, item in enumerate(items, start=2):
        for col_idx, (field, _) in enumerate(EXPORT_COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_excel_cell_value(getattr(item, field)))

    for col_idx, (_, header) in enumerate(EXPORT_COLUMNS, start=1):
        max_length = len(header)
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, values_only=True):
            for value in cell:
                if value is not None:
                    max_length = max(max_length, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@router.get("/kpi", response_model=APIResponse)
async def get_items_kpi(
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """KPI: active, due for review (next_review_date >= today), overdue (next_review_date < today), draft, and by type."""
    data = await items_service.get_items_kpi(db)
    return APIResponse(message="KPI fetched", status_code=200, status="success", data=data)


@router.post("/", response_model=APIResponsePaginated)
async def list_combined(
    body: ItemsListBody | None = Body(None),
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """Paginated list of events and/or documents. All filters and pagination in payload (optional; empty body = defaults)."""
    payload = body or ItemsListBody()
    data, total = await items_service.list_combined_filtered(
        db,
        page=payload.page,
        page_size=payload.page_size,
        item_type=payload.item_type,
        document_types=payload.document_types,
        document_names=payload.document_names,
        statuses=payload.statuses,
        last_updated_start=payload.last_updated_start,
        last_updated_end=payload.last_updated_end,
        next_review_start=payload.next_review_start,
        next_review_end=payload.next_review_end,
        due_for_review=payload.due_for_review,
        overdue=payload.overdue,
        search=payload.search,
    )
    logger.info("list_combined total=%s page=%s", total, payload.page)
    return APIResponsePaginated(
        message="Items fetched",
        status_code=200,
        status="success",
        data=data,
        total=total,
        page=payload.page,
        page_size=payload.page_size,
    )


@router.post("/export")
async def export_combined(
    body: ItemsListBody | None = Body(None),
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    """Export all events/documents matching the selected filters to an Excel file."""
    payload = body or ItemsListBody()
    _, total = await items_service.list_combined_filtered(
        db,
        page=1,
        page_size=1,
        item_type=payload.item_type,
        document_types=payload.document_types,
        document_names=payload.document_names,
        statuses=payload.statuses,
        last_updated_start=payload.last_updated_start,
        last_updated_end=payload.last_updated_end,
        next_review_start=payload.next_review_start,
        next_review_end=payload.next_review_end,
        due_for_review=payload.due_for_review,
        overdue=payload.overdue,
        search=payload.search,
        cache_result=False,
    )
    data, _ = await items_service.list_combined_filtered(
        db,
        page=1,
        page_size=max(total, 1),
        item_type=payload.item_type,
        document_types=payload.document_types,
        document_names=payload.document_names,
        statuses=payload.statuses,
        last_updated_start=payload.last_updated_start,
        last_updated_end=payload.last_updated_end,
        next_review_start=payload.next_review_start,
        next_review_end=payload.next_review_end,
        due_for_review=payload.due_for_review,
        overdue=payload.overdue,
        search=payload.search,
        cache_result=False,
    )
    logger.info("export_combined total=%s", total)

    filename = f"combined-items-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        _build_combined_export_workbook(data),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/{item_id}", response_model=APIResponse)
async def get_item_detail(
    item_id: int,
    item_type: str = Query(..., description=f"'{EVENT}' or '{DOCUMENT}'"),
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    data = await items_service.get_item_detail(db, item_id=item_id, item_type=item_type)
    logger.info("get_item_detail item_id=%s item_type=%s", item_id, item_type)
    return APIResponse(message="Item fetched", status_code=200, status="success", data=data)


@router.get("/{item_id}/revisions", response_model=APIResponse)
async def list_item_revisions(
    item_id: int,
    item_type: str = Query(..., description=f"'{EVENT}' or '{DOCUMENT}'"),
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    data = await items_service.list_item_revisions(db, item_id=item_id, item_type=item_type)
    logger.info("list_item_revisions item_id=%s item_type=%s count=%s", item_id, item_type, len(data))
    return APIResponse(message="Revisions fetched", status_code=200, status="success", data=data)


@router.get("/{item_id}/revisions/{revision_number}", response_model=APIResponse)
async def get_item_revision_snapshot(
    item_id: int,
    revision_number: int,
    item_type: str = Query(..., description=f"'{EVENT}' or '{DOCUMENT}'"),
    db: AsyncSession = Depends(get_db),
    user: CurrentUser = Depends(get_current_user),
):
    data = await items_service.get_item_revision_snapshot(
        db, item_id=item_id, item_type=item_type,
        revision_number=revision_number,
    )
    logger.info(
        "get_item_revision_snapshot item_id=%s item_type=%s rn=%s",
        item_id, item_type, revision_number,
    )
    return APIResponse(message="Revision fetched", status_code=200, status="success", data=data)
